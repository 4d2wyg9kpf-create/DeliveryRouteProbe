from __future__ import annotations

import csv
import dataclasses
import datetime as dt
import json
import sqlite3
import tempfile
import time
import unittest
from contextlib import closing
from pathlib import Path

import openpyxl

from location_analysis.diagnostics import populate, sample_record
from location_analysis.importer import inspect_file, parse_records
from location_analysis.models import split_roads, validate_coordinates
from location_analysis.storage import CatalogStore, write_table


class TemporaryCatalog(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.store = CatalogStore(self.root / "catalog.sqlite3")

    def tearDown(self):
        self.temporary.cleanup()

    def csv_file(self, rows, encoding="utf-8-sig", name="단지정보.csv"):
        path = self.root / name
        with path.open("w", encoding=encoding, newline="") as handle:
            csv.writer(handle).writerows(rows)
        return path


class AddressTests(unittest.TestCase):
    def test_multiple_building_numbers_and_province_alias(self):
        roads, issues = split_roads("대전 서구 문정로 001, 10, 1-1 (탄방동, 예시), 문정로 12")
        self.assertFalse(issues)
        self.assertEqual(roads, ["대전광역시 서구 문정로 1", "대전광역시 서구 문정로 10", "대전광역시 서구 문정로 1-1", "대전광역시 서구 문정로 12"])

    def test_ambiguous_address_is_not_silently_dropped(self):
        roads, issues = split_roads("서울 강서구 예제로 1; 주소확인필요")
        self.assertEqual(len(roads), 1)
        self.assertEqual(issues, ["주소확인필요"])

    def test_full_addresses_and_parenthetical_commas(self):
        roads, issues = split_roads("서울 강서구 예제로 1(화곡동, 가상), 서울 강서구 예제로 2")
        self.assertEqual(len(roads), 2)
        self.assertFalse(issues)

    def test_coordinates_require_both_and_finite(self):
        self.assertEqual(validate_coordinates("36.1", "127.5"), (36.1, 127.5))
        self.assertEqual(validate_coordinates("", ""), (None, None))
        for pair in [("nan", 127), (36, "inf"), (91, 127), (36, ""), (36, 181)]:
            with self.subTest(pair=pair), self.assertRaises(ValueError):
                validate_coordinates(*pair)


class ImportTests(TemporaryCatalog):
    def test_korean_xlsx_title_header_leading_zero_and_date(self):
        path = self.root / "공동주택단지정보.xlsx"
        book = openpyxl.Workbook()
        book.active.title = "빈 시트"
        sheet = book.create_sheet("단지 기본정보")
        sheet.append(["공동주택단지 전국 기본정보"])
        sheet.append(["단지코드", "단지명", "도로명주소", "세대수", "동수", "사용승인일", "주차대수"])
        sheet.append([123, "가상 단지", "서울특별시 강서구 예제로 1", 1234, 12, dt.datetime(2000, 1, 2), 1500])
        sheet.cell(3, 1).number_format = "000000"
        book.save(path)
        info = inspect_file(path)
        self.assertEqual((info.sheet, info.header_row), ("단지 기본정보", 1))
        records, _ = parse_records(info)
        self.assertEqual((records[0].code, records[0].households, records[0].buildings, records[0].approval_date), ("000123", "1234", "12", "2000-01-02"))
        self.assertEqual(records[0].raw["주차대수"], "1500")

    def test_cp949_csv_and_manual_mapping(self):
        path = self.csv_file([["번호", "명칭", "새주소", "기타"], ["A001", "가상 단지", "서울 강서구 예제로 1", "보존"]], "cp949")
        info = inspect_file(path)
        self.assertEqual(info.encoding, "cp949")
        mapping = {"code": "번호", "name": "명칭", "road": "새주소"}
        self.store.import_file(info, mapping=mapping)
        row = self.store.get_complex("A001")
        self.assertEqual(json.loads(row["raw"])["기타"], "보존")

    def test_utf8_sample_ending_inside_multibyte_character(self):
        path = self.csv_file([["단지코드", "단지명", "도로명주소", "기타"], ["A001", "가상", "서울 강서구 예제로 1", "가" * 50000]], encoding="utf-8")
        info = inspect_file(path)
        self.assertEqual(info.encoding, "utf-8-sig")
        self.assertEqual(len(parse_records(info)[0]), 1)

    def test_same_code_multi_address_rows_merge_once(self):
        path = self.csv_file([["단지코드", "단지명", "도로명주소"], ["A001", "가상", "서울 강서구 예제로 1"], ["A001", "가상", "서울 강서구 예제로 2"]])
        result = self.store.import_file(inspect_file(path))
        self.assertEqual((result["total"], result["repeated"]), (1, 1))
        self.assertEqual(self.store.get_complex("A001")["address_count"], 2)
        self.assertEqual(self.store.stats()["candidates"], 1)

    def test_conflicting_same_code_rejects_entire_import(self):
        path = self.csv_file([["단지코드", "단지명", "도로명주소"], ["A001", "단지1", "서울 강서구 예제로 1"], ["A001", "단지2", "서울 강서구 예제로 2"]])
        with self.assertRaises(ValueError):
            self.store.import_file(inspect_file(path))
        self.assertEqual(self.store.stats()["total"], 0)
        self.assertFalse(self.store.imports())

    def test_missing_code_or_road_column_does_not_create_partial_inventory(self):
        path = self.csv_file([["단지코드", "단지명", "도로명주소"], ["A001", "가상", "서울 강서구 예제로 1"], ["", "누락", "서울 강서구 예제로 2"]])
        with self.assertRaises(ValueError):
            self.store.import_file(inspect_file(path))
        path2 = self.csv_file([["단지코드", "단지명"], ["A001", "가상"]], name="주소열없음.csv")
        with self.assertRaises(ValueError):
            self.store.import_file(inspect_file(path2))
        self.assertEqual(self.store.stats()["total"], 0)

    def test_unfetched_legacy_json_is_not_blank_road_exclusion(self):
        path = self.root / "이전목록.json"
        path.write_text(json.dumps({"records": [{"id": "A001", "name": "미수집", "roadAddresses": [], "inventoryProvince": "서울특별시"}, {"id": "A002", "name": "실제공란", "roadAddresses": [], "inventoryProvince": "서울특별시", "fetchedAt": 1234}]}), encoding="utf-8")
        self.store.import_file(inspect_file(path))
        self.assertEqual(self.store.get_complex("A001")["status"], "주소 확인")
        self.assertEqual(self.store.get_complex("A002")["status"], "주소 공란")

    def test_empty_and_unsupported_files(self):
        for name in ["empty.csv", "unknown.pdf"]:
            path = self.root / name
            path.write_text("", encoding="utf-8")
            with self.subTest(name=name), self.assertRaises(ValueError):
                inspect_file(path)


class ReviewTests(TemporaryCatalog):
    def setUp(self):
        super().setUp()
        self.records = populate(self.store)

    def test_same_address_separate_codes_and_multiple_roads_one_candidate(self):
        self.assertEqual(self.store.stats()["candidates"], 4)
        self.assertEqual(self.store.get_complex("TEST-A003")["address_count"], 2)
        rows, count = self.store.list_complexes(page="duplicates")
        self.assertEqual(count, 3)
        self.assertEqual({r["code"] for r in rows}, {"TEST-A001", "TEST-A002", "TEST-A004"})

    def test_many_to_many_links(self):
        self.assertEqual(len(self.store.get_complex("TEST-A003")["links"]), 2)
        self.store.save_review("TEST-A002", verdict="valid", reason="별개 관리단지", selected_road=self.records[1].roads[0], realty_codes=["TEST-P001", "TEST-P002"])
        prices, _ = self.store.list_realty("TEST-P001")
        self.assertEqual(prices[0]["linked_count"], 2)

    def test_exclusion_clears_live_links_and_preserves_raw_and_history(self):
        before = self.store.get_complex("TEST-A001")
        self.store.save_review("TEST-A001", verdict="excluded", reason="중복 확인", latitude="invalid", longitude="invalid")
        after = self.store.get_complex("TEST-A001")
        self.assertEqual(after["status"], "제외")
        self.assertEqual(after["links"], [])
        self.assertIsNone(after["latitude"])
        self.assertEqual(after["raw"], before["raw"])
        self.assertEqual(json.loads(after["history"][0]["before_json"])["links"][0]["realty_code"], "TEST-P001")
        self.store.save_review("TEST-A001", verdict="pending", note="다시 확인")
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "검토 대기")

    def test_confirmation_requires_link_building_membership_and_entrance(self):
        code = "TEST-A006"
        values = dict(verdict="valid", reason="실재 확인")
        self.store.save_review(code, **values)
        self.assertEqual(self.store.get_complex(code)["status"], "대조 대기")
        values.update(realty_codes=["TEST-P003"], selected_road=self.records[5].roads[0])
        self.store.save_review(code, **values)
        self.assertEqual(self.store.get_complex(code)["status"], "건물 확인")
        values.update(building_name="가동", member_confirmed=True)
        self.store.save_review(code, **values)
        self.assertEqual(self.store.get_complex(code)["status"], "입구 확인")
        values.update(latitude=35.1, longitude=129.1, entrance_confirmed=True)
        self.store.save_review(code, **values)
        self.assertEqual(self.store.get_complex(code)["status"], "확정")

    def test_review_validation_and_optimistic_source_check(self):
        invalid = [dict(verdict="excluded"), dict(verdict="valid", reason="확인", selected_road="서울 강서구 다른로 1"), dict(verdict="valid", reason="확인", entrance_confirmed=True), dict(verdict="pending", realty_codes=["NOT-FOUND"]), dict(verdict="pending", expected_scope="outdated")]
        for args in invalid:
            with self.subTest(args=args), self.assertRaises(ValueError):
                self.store.save_review("TEST-A001", **args)
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "확정")

    def test_repeat_import_preserves_review_and_exclusion(self):
        result = self.store.import_records(self.records, source_date="2026-01-02")
        self.assertEqual(result["unchanged"], 6)
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "확정")
        self.assertEqual(self.store.get_complex("TEST-A004")["status"], "제외")

    def test_scope_change_retains_decisions_and_requires_review(self):
        changed = dataclasses.replace(self.records[0], households="777", raw={"수정": "자료"})
        self.store.import_records([changed], source_date="2026-01-03")
        row = self.store.get_complex(changed.code)
        self.assertEqual((row["status"], row["verdict"], row["building_name"]), ("재검토", "valid", "101동"))
        with self.store.connect() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM observations WHERE kind='kapt' AND code=?", (changed.code,)).fetchone()[0], 2)

    def test_linked_realty_change_requires_review(self):
        self.store.import_records([sample_record("TEST-P001", "바뀐 이름", "대전광역시 서구 예제로 10")], kind="realty")
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "재검토")

    def test_raw_only_update_preserves_confirmed_scope(self):
        self.store.import_records([dataclasses.replace(self.records[0], raw={"새로운원본항목": "값"})])
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "확정")

    def test_full_snapshot_missing_is_review_flag_and_can_return(self):
        self.store.import_records([self.records[1]], full_snapshot=True, source_date="2026-02-01")
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "수록 확인")
        self.assertEqual(self.store.stats()["total"], 6)
        self.store.import_records([self.records[0]], source_date="2026-02-02")
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "확정")

    def test_partial_import_does_not_mark_missing(self):
        self.store.import_records([self.records[1]], source_date="2026-02-01")
        self.assertEqual(self.store.get_complex("TEST-A001")["status"], "확정")

    def test_old_or_bad_dates_reject_atomically(self):
        count = len(self.store.imports())
        for date in ["2025-12-31", "2026-02-30"]:
            with self.subTest(date=date), self.assertRaises(ValueError):
                self.store.import_records([sample_record("NEW", "새 단지", "서울 강서구 예제로 1"), self.records[0]], source_date=date, full_snapshot=True)
        self.assertEqual(len(self.store.imports()), count)
        with self.assertRaises(ValueError):
            self.store.get_complex("NEW")

    def test_cancellation_rolls_back_new_and_changed_records(self):
        before = self.store.stats()
        calls = 0
        def cancel():
            nonlocal calls
            calls += 1
            return calls > 1
        with self.assertRaises(InterruptedError):
            self.store.import_records([sample_record("NEW", "새단지", "서울 강서구 예제로 2"), self.records[0]], cancelled=cancel)
        self.assertEqual(self.store.stats(), before)
        with self.assertRaises(ValueError):
            self.store.get_complex("NEW")

    def test_region_search_and_literal_wildcards(self):
        rows, count = self.store.list_complexes(province="대전광역시", district="서구")
        self.assertEqual(count, 3)
        self.assertEqual(self.store.list_complexes(search="%")[1], 0)
        self.assertEqual(self.store.list_complexes(search="' OR 1=1 --")[1], 0)
        self.assertEqual(self.store.list_complexes(search="해오름")[1], 1)

    def test_export_candidate_and_confirmed_counts(self):
        for only, expected in [(False, 4), (True, 1)]:
            path = self.root / f"표준지-{only}.xlsx"
            self.assertEqual(self.store.export_standards(path, confirmed_only=only), expected)
            book = openpyxl.load_workbook(path, read_only=True)
            rows = list(book.active.values)
            book.close()
            self.assertEqual(len(rows), expected + 1)
            self.assertEqual(len({row[0] for row in rows[1:]}), expected)

    def test_export_escapes_formulas(self):
        for ext in ["csv", "xlsx"]:
            path = self.root / ("safe." + ext)
            write_table(path, ["원본"], [["=HYPERLINK(1)"], ["+12"], ["@SUM(A1)"], [12]])
            if ext == "csv":
                with path.open(encoding="utf-8-sig", newline="") as handle:
                    values = list(csv.reader(handle))[1:]
            else:
                book = openpyxl.load_workbook(path, read_only=True)
                values = list(book.active.values)[1:]
                book.close()
            self.assertTrue(all(str(row[0]).startswith("'") for row in values[:3]))

    def test_backup_restore_keeps_raw_links_decisions_history(self):
        backup = self.root / "backup.sqlite3"
        self.store.backup(backup)
        before = self.store.get_complex("TEST-A001")
        self.store.save_review("TEST-A001", verdict="excluded", reason="임시")
        safeguard = self.store.restore(backup)
        self.assertTrue(safeguard.exists())
        after = self.store.get_complex("TEST-A001")
        self.assertEqual(before, after)
        self.assertEqual(CatalogStore(safeguard).get_complex("TEST-A001")["status"], "제외")

    def test_wrong_backup_is_rejected_without_changing_current(self):
        path = self.root / "wrong.sqlite3"
        with closing(sqlite3.connect(path)) as con:
            con.execute("CREATE TABLE other(x TEXT)")
            con.commit()
        before = self.store.stats()
        with self.assertRaises((ValueError, sqlite3.Error)):
            self.store.restore(path)
        self.assertEqual(before, self.store.stats())


class ScaleTests(TemporaryCatalog):
    def test_twenty_thousand_records_import_paginate_and_search(self):
        records = [sample_record(f"SCALE-{i:06}", f"검증 단지 {i:06}", f"경기도 수원시 영통구 가상로 {i+1}") for i in range(20000)]
        start = time.monotonic()
        self.store.import_records(records, full_snapshot=True)
        import_seconds = time.monotonic() - start
        start = time.monotonic()
        rows, total = self.store.list_complexes(offset=19950)
        self.assertEqual((total, len(rows)), (20000, 50))
        self.assertEqual(self.store.list_complexes(search="SCALE-019999")[1], 1)
        self.assertEqual(self.store.stats()["candidates"], 20000)
        self.assertEqual(self.store.regions("경기도")[1], ["수원시 영통구"])
        print(f"SCALE fixture-only rows=20000 import={import_seconds:.3f}s query={time.monotonic()-start:.3f}s")


if __name__ == "__main__":
    unittest.main()
