from __future__ import annotations

import csv
import datetime as dt
import json
import os
import shutil
import sqlite3
import tempfile
from contextlib import closing, contextmanager
from pathlib import Path

from .models import CatalogRecord, clean, fingerprint, json_text, road_key, validate_coordinates
from .importer import FileInfo, file_sha256, parse_records


SCHEMA_VERSION = 1
PAGE_SIZE = 150
STATUSES = ["검토 대기", "주소 확인", "대조 대기", "건물 확인", "입구 확인", "재검토", "확정", "수록 확인", "주소 공란", "제외"]


def now():
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def default_db_path():
    root = Path(os.environ.get("LOCALAPPDATA") or (Path.home() / ".local" / "share"))
    return root / "LocationAnalysisPC" / "catalog.sqlite3"


SCHEMA = """
CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY,value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS imports(
 id INTEGER PRIMARY KEY,kind TEXT NOT NULL,filename TEXT NOT NULL,sha256 TEXT NOT NULL,
 source_date TEXT NOT NULL,imported_at TEXT NOT NULL,full_snapshot INTEGER NOT NULL,
 total INTEGER NOT NULL,added INTEGER NOT NULL,changed INTEGER NOT NULL,unchanged INTEGER NOT NULL,
 blank_roads INTEGER NOT NULL,unparsed_roads INTEGER NOT NULL,repeated INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS complexes(
 code TEXT PRIMARY KEY,name TEXT NOT NULL,road_raw TEXT NOT NULL,roads TEXT NOT NULL,road_issues TEXT NOT NULL,
 jibun TEXT NOT NULL,province TEXT NOT NULL,district TEXT NOT NULL,households TEXT NOT NULL,
 buildings TEXT NOT NULL,use_type TEXT NOT NULL,approval_date TEXT NOT NULL,raw TEXT NOT NULL,
 scope_hash TEXT NOT NULL,import_id INTEGER NOT NULL,source_date TEXT NOT NULL,
 address_count INTEGER NOT NULL,address_issue_count INTEGER NOT NULL,road_known INTEGER NOT NULL DEFAULT 1,missing INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS addresses(
 code TEXT NOT NULL REFERENCES complexes(code),road_key TEXT NOT NULL,PRIMARY KEY(code,road_key));
CREATE INDEX IF NOT EXISTS addresses_road ON addresses(road_key,code);
CREATE INDEX IF NOT EXISTS complexes_region ON complexes(province,district,code);
CREATE TABLE IF NOT EXISTS realty(
 code TEXT PRIMARY KEY,name TEXT NOT NULL,road_raw TEXT NOT NULL,roads TEXT NOT NULL,road_issues TEXT NOT NULL,
 jibun TEXT NOT NULL,province TEXT NOT NULL,district TEXT NOT NULL,households TEXT NOT NULL,
 buildings TEXT NOT NULL,use_type TEXT NOT NULL,approval_date TEXT NOT NULL,raw TEXT NOT NULL,
 scope_hash TEXT NOT NULL,import_id INTEGER NOT NULL,source_date TEXT NOT NULL,
 address_count INTEGER NOT NULL,address_issue_count INTEGER NOT NULL,road_known INTEGER NOT NULL DEFAULT 1,missing INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS realty_addresses(code TEXT NOT NULL REFERENCES realty(code),road_key TEXT NOT NULL,PRIMARY KEY(code,road_key));
CREATE INDEX IF NOT EXISTS realty_addresses_road ON realty_addresses(road_key,code);
CREATE TABLE IF NOT EXISTS reviews(
 code TEXT PRIMARY KEY REFERENCES complexes(code),verdict TEXT NOT NULL CHECK(verdict IN ('pending','valid','excluded')),
 reason TEXT NOT NULL DEFAULT '',note TEXT NOT NULL DEFAULT '',selected_road TEXT NOT NULL DEFAULT '',
 building_name TEXT NOT NULL DEFAULT '',building_id TEXT NOT NULL DEFAULT '',member_confirmed INTEGER NOT NULL DEFAULT 0,
 latitude REAL,longitude REAL,entrance_confirmed INTEGER NOT NULL DEFAULT 0,
 scope_hash TEXT NOT NULL,updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS links(
 kapt_code TEXT NOT NULL REFERENCES complexes(code),realty_code TEXT NOT NULL REFERENCES realty(code),
 realty_scope_hash TEXT NOT NULL,PRIMARY KEY(kapt_code,realty_code));
CREATE INDEX IF NOT EXISTS links_realty ON links(realty_code,kapt_code);
CREATE TABLE IF NOT EXISTS history(
 id INTEGER PRIMARY KEY,at TEXT NOT NULL,action TEXT NOT NULL,code TEXT NOT NULL,before_json TEXT NOT NULL,after_json TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS observations(
 kind TEXT NOT NULL,code TEXT NOT NULL,content_hash TEXT NOT NULL,import_id INTEGER NOT NULL,observed_at TEXT NOT NULL,payload TEXT NOT NULL,
 PRIMARY KEY(kind,code,content_hash));
CREATE VIEW IF NOT EXISTS catalog AS SELECT c.*,r.verdict,r.reason,r.note,r.selected_road,r.building_name,r.building_id,
 r.member_confirmed,r.latitude,r.longitude,r.entrance_confirmed,r.updated_at,
 (SELECT COUNT(*) FROM links l WHERE l.kapt_code=c.code) AS link_count,
 CASE
 WHEN r.verdict='excluded' THEN '제외'
 WHEN c.road_known=0 THEN '주소 확인'
 WHEN c.address_count=0 AND c.address_issue_count=0 THEN '주소 공란'
 WHEN c.missing=1 THEN '수록 확인'
 WHEN c.address_issue_count>0 THEN '주소 확인'
 WHEN r.verdict IS NULL OR r.verdict='pending' THEN '검토 대기'
 WHEN r.scope_hash<>c.scope_hash OR EXISTS(
   SELECT 1 FROM links l JOIN realty p ON p.code=l.realty_code
   WHERE l.kapt_code=c.code AND (l.realty_scope_hash<>p.scope_hash OR p.missing=1)) THEN '재검토'
 WHEN (SELECT COUNT(*) FROM links l WHERE l.kapt_code=c.code)=0 THEN '대조 대기'
 WHEN r.selected_road='' THEN '주소 확인'
 WHEN r.building_name='' OR r.member_confirmed<>1 THEN '건물 확인'
 WHEN r.latitude IS NULL OR r.longitude IS NULL OR r.entrance_confirmed<>1 THEN '입구 확인'
 ELSE '확정' END AS status
 FROM complexes c LEFT JOIN reviews r ON r.code=c.code;
"""


class CatalogStore:
    def __init__(self, path=None):
        self.path = Path(path or default_db_path()).resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as con:
            version = con.execute("PRAGMA user_version").fetchone()[0]
            if version not in {0, SCHEMA_VERSION}:
                raise ValueError("이 보관함은 더 새로운 프로그램에서 만든 파일입니다.")
            con.execute("PRAGMA journal_mode=WAL")
            con.executescript(SCHEMA)
            con.execute(f"PRAGMA user_version={SCHEMA_VERSION}")
            con.execute("INSERT OR REPLACE INTO meta VALUES('application','LocationAnalysisPC')")

    @contextmanager
    def connect(self):
        con = sqlite3.connect(self.path, timeout=30)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA foreign_keys=ON")
        try:
            with con:
                yield con
        finally:
            con.close()

    def import_file(self, info: FileInfo, kind="kapt", mapping=None, source_date="", full_snapshot=False, progress=None, cancelled=None):
        if kind not in {"kapt", "realty"}:
            raise ValueError("지원하지 않는 단지 자료입니다.")
        if source_date:
            try:
                dt.date.fromisoformat(source_date)
            except ValueError:
                raise ValueError("자료 기준일은 YYYY-MM-DD 형식으로 입력하세요.") from None
        selected = mapping or info.mapping
        if kind == "kapt" and not selected.get("road"):
            raise ValueError("K-apt 자료의 도로명주소 열을 연결하세요. 상세 미수집과 실제 공란을 구분하기 위해 필요합니다.")
        records, repeated = parse_records(info, selected, progress, cancelled)
        sha = file_sha256(info.path)
        return self.import_records(records, kind=kind, filename=Path(info.path).name, sha=sha,
            source_date=source_date, full_snapshot=full_snapshot, repeated=repeated, progress=progress, cancelled=cancelled)

    def import_records(self, records: list[CatalogRecord], kind="kapt", filename="직접 등록", sha="", source_date="", full_snapshot=False, repeated=0, progress=None, cancelled=None):
        if kind not in {"kapt", "realty"} or not records:
            raise ValueError("가져올 단지 자료가 없습니다.")
        if source_date:
            try:
                dt.date.fromisoformat(source_date)
            except ValueError:
                raise ValueError("자료 기준일은 YYYY-MM-DD 형식으로 입력하세요.") from None
        table = "complexes" if kind == "kapt" else "realty"
        address_table = "addresses" if kind == "kapt" else "realty_addresses"
        summary = dict(total=len(records), added=0, changed=0, unchanged=0,
            blank_roads=sum(r.road_known and not r.roads and not r.road_issues for r in records),
            unparsed_roads=sum(bool(r.road_issues) for r in records), repeated=repeated)
        codes = [r.code for r in records]
        if len(set(codes)) != len(codes):
            raise ValueError("중복된 단지코드는 파일 가져오기 단계에서 먼저 정리해야 합니다.")
        for record in records:
            record.validate()
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            # An old file cannot silently replace a newer observation of the same source.
            if source_date:
                newest_full = con.execute("SELECT MAX(source_date) FROM imports WHERE kind=? AND full_snapshot=1", (kind,)).fetchone()[0]
                if full_snapshot and newest_full and source_date < newest_full:
                    raise ValueError(f"이미 {newest_full} 기준의 전체 목록이 있습니다. 이전 기준일의 전체 목록은 덮어쓰지 않았습니다.")
            cur = con.execute("INSERT INTO imports(kind,filename,sha256,source_date,imported_at,full_snapshot,total,added,changed,unchanged,blank_roads,unparsed_roads,repeated) VALUES(?,?,?,?,?,?,?,0,0,0,?,?,?)",
                (kind, filename, sha, source_date, now(), int(full_snapshot), len(records), summary["blank_roads"], summary["unparsed_roads"], repeated))
            import_id = cur.lastrowid
            for index, record in enumerate(records, 1):
                if cancelled and cancelled():
                    raise InterruptedError("가져오기를 취소했습니다. 기존 자료는 그대로 유지됩니다.")
                old = con.execute(f"SELECT scope_hash,source_date,raw FROM {table} WHERE code=?", (record.code,)).fetchone()
                if old and source_date and old["source_date"] and source_date < old["source_date"]:
                    raise ValueError(f"{record.code}: 기존 기준일 {old['source_date']}보다 오래된 자료입니다. 기존 자료를 유지했습니다.")
                same = old and old["scope_hash"] == record.scope_hash and old["raw"] == json_text(record.raw)
                summary["unchanged" if same else "changed" if old else "added"] += 1
                values = (record.code,record.name,record.road_raw,json_text(record.roads),json_text(record.road_issues),record.jibun,
                    record.province,record.district,record.households,record.buildings,record.use,record.approval_date,json_text(record.raw),
                    record.scope_hash,import_id,source_date or (old["source_date"] if old else ""),len(record.roads),len(record.road_issues),int(record.road_known))
                columns = "code,name,road_raw,roads,road_issues,jibun,province,district,households,buildings,use_type,approval_date,raw,scope_hash,import_id,source_date,address_count,address_issue_count,road_known"
                updates = ",".join(f"{key}=excluded.{key}" for key in columns.split(",") if key != "code")
                con.execute(f"INSERT INTO {table}({columns}) VALUES({','.join('?' for _ in values)}) ON CONFLICT(code) DO UPDATE SET {updates},missing=0", values)
                payload=json_text(record.__dict__)
                con.execute("INSERT OR IGNORE INTO observations VALUES(?,?,?,?,?,?)",(kind,record.code,fingerprint(record.__dict__),import_id,now(),payload))
                con.execute(f"DELETE FROM {address_table} WHERE code=?", (record.code,))
                con.executemany(f"INSERT INTO {address_table}(code,road_key) VALUES(?,?)", ((record.code, road) for road in record.roads))
                if progress and (index % 200 == 0 or index == len(records)):
                    progress(index, f"{index:,}/{len(records):,}개 단지 저장 준비")
            if full_snapshot:
                # Missing is a review flag, never an automatic demolition/exclusion decision.
                con.execute(f"UPDATE {table} SET missing=1 WHERE import_id<>?", (import_id,))
            con.execute("UPDATE imports SET added=?,changed=?,unchanged=? WHERE id=?", (summary["added"],summary["changed"],summary["unchanged"],import_id))
        return {**summary, "import_id": import_id, "kind": kind}

    def _where(self, page="kapt", search="", province="", district="", status=""):
        conditions, params = [], []
        if page == "standards":
            conditions.append("c.status NOT IN ('제외','주소 공란')")
        elif page == "duplicates":
            conditions.append("EXISTS(SELECT 1 FROM addresses a WHERE a.code=c.code AND a.road_key IN (SELECT x.road_key FROM addresses x JOIN catalog d ON d.code=x.code WHERE d.status<>'제외' GROUP BY x.road_key HAVING COUNT(DISTINCT x.code)>1))")
        if province:
            conditions.append("c.province=?"); params.append(province)
        if district:
            conditions.append("c.district=?"); params.append(district)
        if status:
            conditions.append("c.status=?"); params.append(status)
        if search:
            escaped = search.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            conditions.append("(c.code LIKE ? ESCAPE '\\' OR c.name LIKE ? ESCAPE '\\' OR c.road_raw LIKE ? ESCAPE '\\' OR c.jibun LIKE ? ESCAPE '\\')")
            params.extend([f"%{escaped}%"]*4)
        return (" WHERE " + " AND ".join(conditions) if conditions else ""), params

    def list_complexes(self, page="kapt", search="", province="", district="", status="", offset=0, limit=PAGE_SIZE):
        where, params = self._where(page, search, province, district, status)
        with self.connect() as con:
            total = con.execute("SELECT COUNT(*) FROM catalog c" + where, params).fetchone()[0]
            rows = con.execute("SELECT c.* FROM catalog c" + where + " ORDER BY c.province,c.district,c.name,c.code LIMIT ? OFFSET ?", [*params, limit, offset]).fetchall()
            return [dict(r) for r in rows], total

    def stats(self):
        with self.connect() as con:
            counts = dict(con.execute("SELECT status,COUNT(*) FROM catalog GROUP BY status").fetchall())
            return {"total": sum(counts.values()), "candidates": sum(v for k,v in counts.items() if k not in {"제외","주소 공란"}),
                "confirmed": counts.get("확정",0), "excluded": counts.get("제외",0)+counts.get("주소 공란",0),
                "needs_review": sum(v for k,v in counts.items() if k not in {"제외","주소 공란","확정"}), "statuses": counts,
                "realty": con.execute("SELECT COUNT(*) FROM realty").fetchone()[0]}

    def regions(self, province=""):
        with self.connect() as con:
            provinces = [r[0] for r in con.execute("SELECT DISTINCT province FROM complexes WHERE province<>'' ORDER BY province")]
            districts = [r[0] for r in con.execute("SELECT DISTINCT district FROM complexes WHERE province=? AND district<>'' ORDER BY district", (province,))]
            return provinces, districts

    def get_complex(self, code):
        with self.connect() as con:
            row = con.execute("SELECT * FROM catalog WHERE code=?", (code,)).fetchone()
            if not row:
                raise ValueError("단지를 찾을 수 없습니다.")
            result = dict(row)
            result["links"] = [dict(r) for r in con.execute("SELECT p.code,p.name,p.road_raw,p.jibun,p.scope_hash,l.realty_scope_hash FROM links l JOIN realty p ON p.code=l.realty_code WHERE l.kapt_code=? ORDER BY p.name",(code,))]
            result["history"] = [dict(r) for r in con.execute("SELECT at,action,before_json,after_json FROM history WHERE code=? ORDER BY id DESC LIMIT 30", (code,))]
            result["shared_address"] = [dict(r) for r in con.execute("SELECT DISTINCT c.code,c.name,c.status FROM addresses a JOIN addresses b ON b.road_key=a.road_key JOIN catalog c ON c.code=b.code WHERE a.code=? AND b.code<>a.code ORDER BY c.name", (code,))]
            return result

    def list_realty(self, search="", offset=0, limit=PAGE_SIZE):
        query = " WHERE name LIKE ? OR code LIKE ? OR road_raw LIKE ? OR jibun LIKE ?" if search else ""
        params = [f"%{search}%"]*4 if search else []
        with self.connect() as con:
            total = con.execute("SELECT COUNT(*) FROM realty"+query,params).fetchone()[0]
            rows = con.execute("SELECT p.*,(SELECT COUNT(*) FROM links l WHERE l.realty_code=p.code) AS linked_count FROM realty p"+query+" ORDER BY name,code LIMIT ? OFFSET ?",[*params,limit,offset]).fetchall()
            return [dict(r) for r in rows],total

    def save_review(self, code, *, verdict, reason="", note="", selected_road="", building_name="", building_id="", member_confirmed=False,
            latitude=None, longitude=None, entrance_confirmed=False, realty_codes=(), expected_scope=None):
        if verdict not in {"pending","valid","excluded"}:
            raise ValueError("검토 상태를 확인하세요.")
        reason, note, building_name, building_id = map(clean,(reason,note,building_name,building_id))
        if verdict in {"valid","excluded"} and not (reason or note):
            raise ValueError("확인 또는 제외한 근거를 입력하세요.")
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            source = con.execute("SELECT * FROM complexes WHERE code=?",(code,)).fetchone()
            if not source:
                raise ValueError("단지를 찾을 수 없습니다.")
            if expected_scope is not None and expected_scope != source["scope_hash"]:
                raise ValueError("검토 중 원본이 변경됐습니다. 창을 다시 열어 확인하세요.")
            before = self._review_snapshot(con,code)
            roads = json.loads(source["roads"])
            if verdict == "valid" and not roads:
                raise ValueError("도로명주소를 확인한 뒤 유효 단지로 지정하세요.")
            if selected_road and selected_road not in roads:
                raise ValueError("대표 주소는 해당 단지의 도로명주소 중에서 선택하세요.")
            lat,lon = (None,None) if verdict=='excluded' else validate_coordinates(latitude,longitude)
            if verdict!='excluded' and entrance_confirmed and (lat is None or lon is None or not building_name or not member_confirmed):
                raise ValueError("소속 건물·동 명칭과 위도·경도를 확인한 뒤 입구 확인을 표시하세요.")
            targets = list(dict.fromkeys(realty_codes))
            if verdict == "excluded":
                selected_road,building_name,building_id,member_confirmed,lat,lon,entrance_confirmed,targets = "","","",False,None,None,False,[]
            target_hashes = []
            for target in targets:
                row = con.execute("SELECT scope_hash FROM realty WHERE code=?",(target,)).fetchone()
                if not row:
                    raise ValueError(f"연결할 공시가격 단지 {target}를 찾을 수 없습니다.")
                target_hashes.append((code,target,row[0]))
            con.execute("INSERT INTO reviews VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(code) DO UPDATE SET verdict=excluded.verdict,reason=excluded.reason,note=excluded.note,selected_road=excluded.selected_road,building_name=excluded.building_name,building_id=excluded.building_id,member_confirmed=excluded.member_confirmed,latitude=excluded.latitude,longitude=excluded.longitude,entrance_confirmed=excluded.entrance_confirmed,scope_hash=excluded.scope_hash,updated_at=excluded.updated_at",
                (code,verdict,reason,note,selected_road,building_name,building_id,int(member_confirmed),lat,lon,int(entrance_confirmed),source["scope_hash"],now()))
            con.execute("DELETE FROM links WHERE kapt_code=?",(code,))
            con.executemany("INSERT INTO links VALUES(?,?,?)",target_hashes)
            after = self._review_snapshot(con,code)
            con.execute("INSERT INTO history(at,action,code,before_json,after_json) VALUES(?,?,?,?,?)",(now(),"검토 저장",code,json_text(before),json_text(after)))

    @staticmethod
    def _review_snapshot(con, code):
        row = con.execute("SELECT * FROM reviews WHERE code=?",(code,)).fetchone()
        return {"review": dict(row) if row else None,"links": [dict(r) for r in con.execute("SELECT * FROM links WHERE kapt_code=? ORDER BY realty_code",(code,))]}

    def imports(self):
        with self.connect() as con:
            return [dict(r) for r in con.execute("SELECT * FROM imports ORDER BY id DESC")]

    def export_standards(self, path, confirmed_only=False):
        rows,_ = self.list_complexes(page="standards",status="확정" if confirmed_only else "",limit=-1)
        headers = ["표준지ID","K-apt단지코드","단지명","검토상태","시도","시군구","대표도로명주소","전체도로명주소","지번주소","표준건물동명","건축물대장ID","입구위도","입구경도","공시가격연결ID","제외확인근거","검토메모","자료기준일"]
        result = []
        with self.connect() as con:
            link_map = {}
            for r in con.execute("SELECT kapt_code,realty_code FROM links ORDER BY realty_code"):
                link_map.setdefault(r[0],[]).append(r[1])
        for r in rows:
            roads = json.loads(r["roads"])
            # Single-address candidates may be displayed, but never called confirmed without all gates.
            result.append(["kapt:"+r["code"],r["code"],r["name"],r["status"],r["province"],r["district"],r["selected_road"] or (roads[0] if len(roads)==1 else ""),
                "\n".join(roads),r["jibun"],r["building_name"] or "",r["building_id"] or "",r["latitude"],r["longitude"],"\n".join(link_map.get(r["code"],[])),r["reason"] or "",r["note"] or "",r["source_date"]])
        write_table(path,headers,result)
        return len(rows)

    def backup(self, destination):
        destination = Path(destination).resolve()
        if destination == self.path:
            raise ValueError("현재 보관함과 다른 파일명을 지정하세요.")
        destination.parent.mkdir(parents=True,exist_ok=True)
        handle, temporary = tempfile.mkstemp(prefix=".catalog-",suffix=".sqlite3",dir=destination.parent)
        os.close(handle)
        try:
            with self.connect() as source, closing(sqlite3.connect(temporary)) as target:
                source.backup(target)
            os.replace(temporary,destination)
        finally:
            Path(temporary).unlink(missing_ok=True)

    @staticmethod
    def validate_backup(path):
        source = Path(path).resolve()
        con = sqlite3.connect(source.as_uri()+"?mode=ro", uri=True)
        try:
            if con.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise ValueError("보관함 무결성 검사에 실패했습니다.")
            if con.execute("PRAGMA user_version").fetchone()[0] != SCHEMA_VERSION or con.execute("SELECT value FROM meta WHERE key='application'").fetchone()[0] != "LocationAnalysisPC":
                raise ValueError("지원하지 않는 보관함 파일입니다.")
            con.execute("SELECT COUNT(*) FROM catalog").fetchone()
            if con.execute("PRAGMA foreign_key_check").fetchall():
                raise ValueError("보관함 연결 관계가 올바르지 않습니다.")
        finally:
            con.close()

    def restore(self, source):
        source = Path(source).resolve()
        if source == self.path:
            raise ValueError("현재 보관함과 다른 파일을 선택하세요.")
        self.validate_backup(source)
        safeguard = self.path.with_name("catalog-before-restore-"+dt.datetime.now().strftime("%Y%m%d-%H%M%S-%f")+".sqlite3")
        self.backup(safeguard)
        # No persistent connection is kept by the UI. Backup writes all pages atomically under SQLite's transaction.
        remote = sqlite3.connect(source.as_uri()+"?mode=ro",uri=True)
        try:
            with self.connect() as target:
                remote.backup(target)
        finally:
            remote.close()
        return safeguard


def spreadsheet_safe(value):
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'"+value
    return value


def write_table(path, headers, rows):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        book = openpyxl.Workbook()
        sheet = book.active; sheet.title = "표준지 목록"
        sheet.append(headers)
        for row in rows:
            sheet.append([spreadsheet_safe(x) for x in row])
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for c in sheet[1]:
            c.fill = PatternFill("solid",fgColor="16394A"); c.font=Font(color="FFFFFF",bold=True)
        for column in sheet.columns:
            letter=column[0].column_letter
            sheet.column_dimensions[letter].width = min(55,max(16,len(str(column[0].value))+5))
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(vertical="top",wrap_text=True)
        book.save(path)
    else:
        with path.open("w",encoding="utf-8-sig",newline="") as handle:
            writer=csv.writer(handle); writer.writerow(headers)
            writer.writerows([[spreadsheet_safe(v) for v in row] for row in rows])
